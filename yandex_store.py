"""Яндекс.Диск — источник правды вместо Google Sheets.

Мастер-файл: disk:/Лагерь/Клиенты/clients.xlsx (вся книга: Общая, сезоны,
вольные листы). Перед каждой заливкой текущий remote уходит в бэкап
(disk:/Лагерь/Бэкапы, держим 10 штук) — last-write-wins, но ничего
не теряется. Локальный clients.xlsx — кэш.
Ни одна функция не роняет вызывающий код: без токена всё молча
пропускается (print в лог).
"""
import os
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
LOCAL_XLSX = os.path.join(BASE, 'clients.xlsx')
SYNC_FILE = os.path.join(BASE, '.ya_sync')

MASTER_PATH = 'disk:/Лагерь/Клиенты/clients.xlsx'
BACKUP_DIR = 'disk:/Лагерь/Бэкапы'
SNAP_DIR = os.path.join(BASE, 'snapshots')
MAX_BACKUPS = 10

# Отложенная заливка: bulk-операции ставят defer(True), делают N правок
# без сети, потом defer(False) + одна заливка. Счётчик — вложенность безопасна.
_DEFER = 0
LAST_FILE = os.path.join(BASE, '.ya_last')


def get_token() -> str:
    return os.environ.get('YANDEX_DISK_TOKEN', '').strip()


def status() -> dict:
    return {'token': bool(get_token())}


def _client():
    tok = get_token()
    if not tok:
        return None
    import yadisk
    return yadisk.YaDisk(token=tok)


def _ensure_dirs(y) -> None:
    for d in ('disk:/Лагерь', 'disk:/Лагерь/Клиенты', BACKUP_DIR):
        try:
            if not y.exists(d):
                y.mkdir(d)
        except Exception:
            pass


def _mark_sync() -> None:
    try:
        with open(SYNC_FILE, 'w', encoding='utf-8') as f:
            f.write(datetime.now().isoformat())
    except Exception:
        pass


def last_sync() -> str:
    try:
        return open(SYNC_FILE, encoding='utf-8').read().strip()
    except Exception:
        return ''


def remote_exists() -> bool:
    try:
        y = _client()
        return bool(y) and bool(y.exists(MASTER_PATH))
    except Exception as e:
        print(f'[Yandex] remote_exists: {e}')
        return False


def download_master() -> dict:
    """Скачать мастер-файл с Диска в локальный clients.xlsx."""
    y = _client()
    if y is None:
        return {'ok': False, 'reason': 'нет YANDEX_DISK_TOKEN'}
    try:
        _ensure_dirs(y)
        if not y.exists(MASTER_PATH):
            return {'ok': False, 'reason': 'мастер-файла нет на Диске'}
        y.download(MASTER_PATH, LOCAL_XLSX)
        _mark_sync()
        return {'ok': True, 'reason': 'скачано с Яндекс.Диска'}
    except Exception as e:
        if os.path.exists(LOCAL_XLSX):
            return {'ok': True, 'rows': -1, 'sheets': {},
                    'reason': f'Диск недоступен ({e}), работаю локально'}
        return {'ok': False, 'reason': f'ошибка скачивания: {e}'}


def backup_remote() -> dict:
    """Копия текущего мастера в Бэкапы + чистка старых (оставить 10)."""
    y = _client()
    if y is None:
        return {'ok': False, 'reason': 'нет токена'}
    try:
        _ensure_dirs(y)
        if not y.exists(MASTER_PATH):
            return {'ok': True, 'reason': 'бэкапить нечего'}
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        dst = f'{BACKUP_DIR}/clients_{stamp}.xlsx'
        y.copy(MASTER_PATH, dst, overwrite=True)
        try:
            items = sorted(
                [p['path'] for p in y.listdir(BACKUP_DIR)
                 if p['path'].endswith('.xlsx')])
            for old in items[:-MAX_BACKUPS]:
                try:
                    y.remove(old, permanently=False)
                except Exception:
                    pass
        except Exception:
            pass
        return {'ok': True, 'reason': dst}
    except Exception as e:
        return {'ok': False, 'reason': f'бэкап не удался: {e}'}


def upload_master(reason: str = '') -> dict:
    """Залить локальный xlsx как мастер (сначала бэкап remote)."""
    y = _client()
    if y is None:
        return {'ok': False, 'reason': 'нет токена, только локально'}
    if not os.path.exists(LOCAL_XLSX):
        return {'ok': False, 'reason': 'нет локального файла'}
    try:
        _ensure_dirs(y)
        br = backup_remote()
        if not br.get('ok'):
            return {'ok': False,
                    'reason': 'мастер НЕ тронут: ' + br.get('reason', '')}
        y.upload(LOCAL_XLSX, MASTER_PATH, overwrite=True)
        _mark_sync()
        return {'ok': True, 'reason': f'залито на Диск ({reason})'}
    except Exception as e:
        return {'ok': False, 'reason': f'заливка не удалась: {e}'}


def sync_after_change(reason: str = '') -> dict:
    """Вызывать из storage после КАЖДОЙ мутации. Не роняет."""
    if _DEFER > 0:
        r = {'ok': None, 'reason': 'идёт bulk-заливка…'}
        _save_last(r)
        return r
    try:
        r = upload_master(reason)
        _save_last(r)
        if not r.get('ok'):
            print(f"[Yandex] {r.get('reason')}")
        return r
    except Exception as e:
        print(f'[Yandex] только локально ({e})')
        return {'ok': False, 'reason': str(e)}


def defer(on: bool) -> dict:
    """Вкл/выкл отложенной заливки (счётчик — вложенность безопасна).
    Обнуление счётчика = одна заливка."""
    global _DEFER
    if on:
        _DEFER += 1
        return {'ok': True, 'reason': 'defer on'}
    _DEFER = max(0, _DEFER - 1)
    if _DEFER == 0:
        return sync_after_change('bulk')
    return {'ok': True, 'reason': 'defer вложен'}


def _save_last(r: dict) -> None:
    try:
        import json
        ok = r.get('ok')
        with open(LAST_FILE, 'w', encoding='utf-8') as f:
            json.dump({'ok': (None if ok is None else bool(ok)),
                       'reason': str(r.get('reason', ''))[:160],
                       'ts': datetime.now().isoformat()}, f,
                      ensure_ascii=False)
    except Exception:
        pass


def last_status() -> dict:
    """Последний результат заливки (для индикатора в UI)."""
    try:
        import json
        return json.load(open(LAST_FILE, encoding='utf-8'))
    except Exception:
        return {'ok': None, 'reason': 'ещё не сохранялось', 'ts': ''}


def snapshot(reason: str = '') -> dict:
    """Снапшот текущего xlsx: локально + копия в Бэкапы на Диске."""
    import shutil
    if not os.path.exists(LOCAL_XLSX):
        return {'ok': False, 'reason': 'нет локального файла'}
    try:
        os.makedirs(SNAP_DIR, exist_ok=True)
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        name = f'clients_{stamp}.xlsx'
        shutil.copy(LOCAL_XLSX, os.path.join(SNAP_DIR, name))
        keep = sorted(os.listdir(SNAP_DIR))[:-MAX_BACKUPS]
        for old in keep:
            try:
                os.remove(os.path.join(SNAP_DIR, old))
            except Exception:
                pass
    except Exception as e:
        return {'ok': False, 'reason': f'локальный снапшот: {e}'}
    y = _client()
    if y is None:
        return {'ok': True, 'reason': f'локально {name}, на Диск не ушло (нет токена)'}
    try:
        _ensure_dirs(y)
        y.upload(os.path.join(SNAP_DIR, name), f'{BACKUP_DIR}/{name}',
                 overwrite=True)
        return {'ok': True, 'reason': name}
    except Exception as e:
        return {'ok': True, 'reason': f'локально {name}, Диск: {e}'}


def list_backups() -> list:
    """Список бэкапов/снапшотов на Диске (новые сверху)."""
    y = _client()
    if y is None:
        return []
    try:
        _ensure_dirs(y)
        items = [p['path'] for p in y.listdir(BACKUP_DIR)
                 if p['path'].endswith('.xlsx')]
        return sorted(items, reverse=True)
    except Exception as e:
        print(f'[Yandex] list_backups: {e}')
        return []


def restore_backup(path: str) -> dict:
    """Восстановить мастер из бэкапа: снапшот текущего → скачать бэкап → залить как мастер."""
    y = _client()
    if y is None:
        return {'ok': False, 'reason': 'нет токена'}
    try:
        snapshot('перед откатом')
        y.download(path, LOCAL_XLSX)
        _mark_sync()
        up = upload_master('откат к ' + path.split('/')[-1])
        if up.get('ok'):
            return {'ok': True, 'reason': 'откатился к ' + path.split('/')[-1]}
        return {'ok': False, 'reason': up.get('reason')}
    except Exception as e:
        return {'ok': False, 'reason': f'не восстановилось: {e}'}


def migrate_from_google() -> dict:
    """Разовый переезд: все листы Google Sheets → локальный xlsx → мастер на Диске.

    Требует оба токена (на Render они оба в env). Google после этого
    не трогается.
    """
    import storage
    import google_sync

    sid = google_sync.get_sheet_id()
    if not sid:
        return {'ok': False, 'reason': 'нет GOOGLE_SHEET_ID'}
    try:
        svc = google_sync._service('sheets', 'v4')
    except Exception as e:
        return {'ok': False, 'reason': f'Google недоступен: {e}'}
    try:
        from openpyxl import Workbook
        wb = Workbook()
        first = True
        per, total = {}, 0
        with google_sync._API_LOCK:
            titles = google_sync.list_sheets(sid, _svc=svc) or list(storage.SHEETS)
            for s in titles:
                try:
                    res = svc.spreadsheets().values().get(
                        spreadsheetId=sid, range=f'{s}!A:ZZ').execute()
                    values = res.get('values', [])
                except Exception:
                    values = []
                if not values:
                    header, data_rows = [h for _, h in storage.CORE], []
                else:
                    header = [str(x or '').strip() or f'col{i+1}'
                              for i, x in enumerate(values[0])]
                    data_rows = [r for r in values[1:] if any(r)]
                ws = wb.active if first else (
                    wb[s] if s in wb.sheetnames else wb.create_sheet(s))
                if first:
                    ws.title = s
                    first = False
                else:
                    if ws.max_row >= 1:
                        ws.delete_rows(1, ws.max_row)
                        ws.delete_cols(1, ws.max_column)
                ws.append(header or [h for _, h in storage.CORE])
                width = len(ws[1])
                for r in data_rows:
                    ws.append(list(r) + [''] * max(0, width - len(r)))
                per[s] = len(data_rows)
                total += len(data_rows)
        for s in storage.SHEETS:
            if s not in wb.sheetnames:
                wb.create_sheet(s).append([h for _, h in storage.CORE])
        wb.save(LOCAL_XLSX)
    except Exception as e:
        return {'ok': False, 'reason': f'не собрал xlsx: {e}'}
    up = upload_master('переезд с Google')
    if not up.get('ok'):
        return {'ok': False, 'reason': up.get('reason')}
    detail = ', '.join(f'{k}: {v}' for k, v in per.items())
    return {'ok': True, 'rows': total, 'sheets': per,
            'reason': f'переехало с Google: всего {total} строк. {detail}'}
